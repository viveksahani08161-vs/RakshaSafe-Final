# -*- coding: utf-8 -*-
"""Insert Figure 4.4 (ER) and Figures 4.5-4.7 (DFD L0/L1/L2) into Chapter 4,
then regenerate the REAL interface build (30 images + these 4 new diagrams).

Approach: perform insertion on the base document, then run the same replacement
logic used for the REAL screenshots build (ch4_real_build.py) on the new base.
"""
import os, re, shutil, copy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = r'C:\Users\SAI\Desktop\Raksha\documentation\CHAPTER_4_RakshaSafe_System_Design.docx'
WORK = r'C:\Users\SAI\AppData\Local\Temp\opencode\ch4_base_with_er_dfd.docx'
OUT = r'C:\Users\SAI\Desktop\Raksha\documentation\CHAPTER_4_RakshaSafe_System_Design_UML_REAL.docx'
BW = r'C:\Users\SAI\Desktop\Raksha\documentation\ch4_bw'
REAL = r'C:\Users\SAI\Desktop\Raksha\documentation\test-evidence\screenshots'

shutil.copy(SRC, WORK)
doc = Document(WORK)

def insert_image_block(doc, before_heading, img_path, caption, width_cm):
    """Insert [IMAGE + caption] paragraphs immediately before the heading."""
    target = None
    for p in doc.paragraphs:
        if p.text.strip() == before_heading:
            target = p
            break
    if target is None:
        print('WARN heading not found:', before_heading)
        return False
    # caption paragraph
    cap = target.insert_paragraph_before()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Cm(0.3)
    r = cap.add_run(caption)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r.italic = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    # image paragraph (so caption comes first, then image is above it in flow? actually
    # insert_paragraph_before appends just before target; we want image above caption)
    imgp = cap.insert_paragraph_before()
    imgp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    imgp.paragraph_format.space_before = Cm(0.2)
    run = imgp.add_run()
    run.add_picture(img_path, width=Cm(width_cm))
    print('inserted:', caption)
    return True

# ---- Figure 4.4 ER before 4.1.4 API Design
insert_image_block(doc, '4.1.4 API Design',
                   os.path.join(BW, 'bw_4_4_er_diagram.png'),
                   'Figure 4.4: RakshaSafe Entity-Relationship Diagram (MongoDB Collections)', 16.5)
# ---- Figures 4.5-4.7 DFD before 4.1.2 Technology Used
insert_image_block(doc, '4.1.2 Technology Used',
                   os.path.join(BW, 'bw_4_5_dfd_level0.png'),
                   'Figure 4.5: Data Flow Diagram - Level 0 (Context Diagram)', 13.5)
insert_image_block(doc, '4.1.2 Technology Used',
                   os.path.join(BW, 'bw_4_6_dfd_level1.png'),
                   'Figure 4.6: Data Flow Diagram - Level 1 (Process Decomposition)', 16.5)
insert_image_block(doc, '4.1.2 Technology Used',
                   os.path.join(BW, 'bw_4_7_dfd_level2.png'),
                   'Figure 4.7: Data Flow Diagram - Level 2 (SOS / Incident Processing)', 14.5)

doc.save(WORK)
print('saved working base with ER + DFD inserted')

# =============================================================
# Now apply the REAL screenshot replacement (same logic as before)
# =============================================================
SLUG_MAP = {
    1:  ('bw_4_2_1_login.png',            '01_login_page.png'),
    2:  ('bw_4_2_2_registration.png',     '03_register_form_filled.png'),
    3:  ('bw_4_2_3_user_dashboard.png',   '04_user_dashboard_after_register.png'),
    4:  ('bw_4_2_4_sos.png',              '06_sos_form_filled.png'),
    5:  ('bw_4_2_5_incident_detail.png',  '27_user_incident_detail.png'),
    6:  ('bw_4_2_6_contacts.png',         '13_emergency_contacts_page.png'),
    7:  ('bw_4_2_7_resources.png',        '12_resources_page.png'),
    8:  ('bw_4_2_8_unsafe.png',           '25_user_unsafe_report.png'),
    9:  ('bw_4_2_9_risk.png',             '26_risk_assessment_admin.png'),
    10: ('bw_4_2_10_notifications.png',   '11_notifications_page.png'),
    11: ('bw_4_2_11_profile.png',         '14_profile_page.png'),
    12: ('bw_4_2_12_admin_dashboard.png', '17_admin_dashboard.png'),
    13: ('bw_4_2_13_admin_incidents.png', '18_admin_incidents.png'),
    14: ('bw_4_2_14_admin_masters.png',   '20_admin_facilities.png'),
}

doc2 = Document(WORK)

def replace_picture_in_paragraph(p, img_path, width_cm):
    for el in p._p.findall(qn('w:r')):
        if el.find(qn('w:drawing')) is not None:
            p._p.remove(el)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(0.2)
    p.paragraph_format.space_after = Cm(0.2)
    run = p.add_run()
    run.add_picture(img_path, width=Cm(width_cm))

replaced_bw = 0
replaced_real = 0
paras = doc2.paragraphs
i = 0
while i < len(paras):
    p = paras[i]
    t = p.text.strip()
    if t.startswith('Figure 4.1'):
        if i > 0 and ('graphic' in paras[i-1]._p.xml or 'blip' in paras[i-1]._p.xml):
            replace_picture_in_paragraph(paras[i-1], os.path.join(BW, 'bw_4_1_architecture.png'), 16.0)
            replaced_bw += 1
    if t.startswith('Figure 4.3'):
        if i > 0 and ('graphic' in paras[i-1]._p.xml or 'blip' in paras[i-1]._p.xml):
            replace_picture_in_paragraph(paras[i-1], os.path.join(BW, 'bw_4_3_deployment.png'), 15.0)
            replaced_bw += 1
    mflow = re.match(r'^Flowchart 4\.2\.(\d+):', t)
    mfig = re.match(r'^Figure 4\.2\.(\d+):', t)
    if mflow:
        num = int(mflow.group(1))
        if num in SLUG_MAP:
            bw_name, _ = SLUG_MAP[num]
            if i > 0 and ('graphic' in paras[i-1]._p.xml or 'blip' in paras[i-1]._p.xml):
                replace_picture_in_paragraph(paras[i-1], os.path.join(BW, bw_name), 12.8)
                replaced_bw += 1
    if mfig:
        num = int(mfig.group(1))
        if num in SLUG_MAP:
            _, real_name = SLUG_MAP[num]
            real_img = os.path.join(REAL, real_name)
            if i > 0 and ('graphic' in paras[i-1]._p.xml or 'blip' in paras[i-1]._p.xml):
                replace_picture_in_paragraph(paras[i-1], real_img, 15.0)
                replaced_real += 1
    i += 1

print('replaced: bw=%d real=%d' % (replaced_bw, replaced_real))
print('inline images:', len(doc2.inline_shapes))
doc2.save(OUT)
print('SAVED:', OUT, round(os.path.getsize(OUT) / 1024), 'KB')