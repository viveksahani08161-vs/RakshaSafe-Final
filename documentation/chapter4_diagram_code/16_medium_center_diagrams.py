# -*- coding: utf-8 -*-
"""Set the 16 user flow diagrams to MEDIUM width, dead-center on page.
  14 flowcharts -> 12.5 cm | architecture -> 14.0 cm | deployment -> 13.5 cm
Screenshots untouched.
"""
import os, io
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image

DOC = r'C:\Users\SAI\Desktop\Raksha\documentation\CHAPTER_4_RakshaSafe_System_Design_UML_REAL_CENTERED.docx'
UBW = r'C:\Users\SAI\Desktop\Raksha\documentation\ch4_bw_user'

SIZE2FILE = {
    (2199, 2167): ('ubw_4_2_1_login.png', 12.5),
    (2027, 1644): ('ubw_4_2_2_registration.png', 12.5),
    (1926, 1434): ('ubw_4_2_3_user_dashboard.png', 12.5),
    (2067, 2922): ('ubw_4_2_4_sos.png', 12.5),
    (2017, 2167): ('ubw_4_2_5_incident_detail.png', 12.5),
    (1935, 2038): ('ubw_4_2_6_contacts.png', 12.5),
    (1932, 1604): ('ubw_4_2_7_resources.png', 12.5),
    (2155, 1948): ('ubw_4_2_8_unsafe.png', 12.5),
    (2098, 1988): ('ubw_4_2_9_risk.png', 12.5),
    (1866, 1614): ('ubw_4_2_10_notifications.png', 12.5),
    (1857, 1524): ('ubw_4_2_11_profile.png', 12.5),
    (1956, 1390): ('ubw_4_2_12_admin_dashboard.png', 12.5),
    (2090, 2088): ('ubw_4_2_13_admin_incidents.png', 12.5),
    (1997, 1610): ('ubw_4_2_14_admin_masters.png', 12.5),
    (2212, 1625): ('ubw_4_1_architecture.png', 14.0),
    (2729, 1193): ('ubw_4_3_deployment.png', 13.5),
}

def para_has_image(p):
    return ('graphic' in p._p.xml) or ('blip' in p._p.xml)

def embedded_size(p, doc):
    for blip in p._p.findall('.//' + qn('a:blip')):
        rid = blip.get(qn('r:embed'))
        if rid in doc.part.rels:
            try:
                return Image.open(io.BytesIO(doc.part.rels[rid].target_part.blob)).size
            except Exception:
                return None
    return None

doc = Document(DOC)
fixed = 0
for p in doc.paragraphs:
    if not para_has_image(p):
        continue
    sz = embedded_size(p, doc)
    if sz in SIZE2FILE:
        fname, wcm = SIZE2FILE[sz]
        for el in p._p.findall(qn('w:r')):
            if el.find(qn('w:drawing')) is not None:
                p._p.remove(el)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Cm(0.3)
        p.paragraph_format.space_after = Cm(0.3)
        run = p.add_run()
        run.add_picture(os.path.join(UBW, fname), width=Cm(wcm))
        fixed += 1
        print('medium centered [%s cm]: %s' % (wcm, fname))

print('fixed:', fixed, '| images:', len(doc.inline_shapes))
doc.save(DOC)
print('SAVED:', DOC, round(os.path.getsize(DOC) / 1024), 'KB')