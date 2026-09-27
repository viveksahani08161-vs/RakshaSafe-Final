"""Fill Chapter 4 System Design docx images + insert architecture/deployment diagrams."""
import os, re, shutil, sys
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

src = r'C:\Users\SAI\Desktop\Raksha\documentation\CHAPTER_4_RakshaSafe_System_Design.docx'
out = r'C:\Users\SAI\Desktop\Raksha\documentation\CHAPTER_4_RakshaSafe_System_Design_FINAL.docx'
imgdir = r'C:\Users\SAI\Desktop\Raksha\documentation\ch4_images'
bak = r'C:\Users\SAI\AppData\Local\Temp\opencode\CHAPTER_4_RakshaSafe_System_Design_preimg.docx'
shutil.copy(src, bak)
print('backup ->', bak)

doc = Document(src)

SLUG = {1: 'login', 2: 'registration', 3: 'user_dashboard', 4: 'sos', 5: 'incident_detail',
        6: 'contacts', 7: 'resources', 8: 'unsafe', 9: 'risk', 10: 'notifications',
        11: 'profile', 12: 'admin_dashboard', 13: 'admin_incidents', 14: 'admin_masters'}

replaced = 0
for p in doc.paragraphs:
    t = p.text.strip()
    m = re.match(r'^\[(Flowchart|Figure) (4\.2\.(\d+)):', t)
    if not m:
        continue
    kind, label, num = m.group(1), m.group(2), int(m.group(3))
    slug = SLUG.get(num)
    if not slug:
        continue
    if kind == 'Flowchart':
        img = os.path.join(imgdir, f'flow_4_2_{num}_{slug}.png')
        width = Cm(12.8)
    else:
        img = os.path.join(imgdir, f'shot_4_2_{num}_{slug}.png')
        width = Cm(15.0)
    if not os.path.exists(img):
        print('MISSING', img)
        continue
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(0.2)
    p.paragraph_format.space_after = Cm(0.2)
    r = p.add_run()
    r.add_picture(img, width=width)
    replaced += 1
print('placeholder images embedded:', replaced)

# ---- insert architecture diagram before heading '4.1.1 Overall System Working'
def insert_image_before(doc, heading_text, img_rel, width_cm, caption):
    img = os.path.join(imgdir, img_rel)
    for p in doc.paragraphs:
        if p.text.strip() == heading_text:
            newp = p.insert_paragraph_before()
            newp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            newp.paragraph_format.space_before = Cm(0.2)
            r = newp.add_run()
            r.add_picture(img, width=Cm(width_cm))
            cap = p.insert_paragraph_before()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.add_run(caption).bold = True
            cap.paragraph_format.space_after = Cm(0.2)
            print('inserted diagram:', caption)
            return True
    print('WARN heading not found:', heading_text)
    return False

insert_image_before(doc, '4.1.1 Overall System Working', 'diagram_4_1_architecture.png', 16.0,
                    'Figure 4.1: RakshaSafe System Architecture')
insert_image_before(doc, '4.3.2 Security Design', 'diagram_4_3_deployment.png', 15.0,
                    'Figure 4.3: RakshaSafe Deployment View')

doc.save(out)
import os as _os
print('saved:', out, round(_os.path.getsize(out) / 1024), 'KB')