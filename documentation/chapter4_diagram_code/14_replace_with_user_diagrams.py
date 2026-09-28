# -*- coding: utf-8 -*-
"""Replace the 16 matplotlib B&W diagrams in Chapter 4 UML_REAL with the
user-provided diagrams (converted to B&W), embedded BIGGER:
  flowcharts  12.8cm -> 15.0cm
  architecture 16.0cm -> 16.5cm
  deployment   15.0cm -> 16.0cm
Everything else (14 real screenshots, ER, DFD L0/L1/L2, text) stays untouched.
"""
import os, re, io
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image

DOC = r'C:\Users\SAI\Desktop\Raksha\documentation\CHAPTER_4_RakshaSafe_System_Design_UML_REAL.docx'
UBW = r'C:\Users\SAI\Desktop\Raksha\documentation\ch4_bw_user'

FLOW_MAP = {
    1: 'ubw_4_2_1_login.png',
    2: 'ubw_4_2_2_registration.png',
    3: 'ubw_4_2_3_user_dashboard.png',
    4: 'ubw_4_2_4_sos.png',
    5: 'ubw_4_2_5_incident_detail.png',
    6: 'ubw_4_2_6_contacts.png',
    7: 'ubw_4_2_7_resources.png',
    8: 'ubw_4_2_8_unsafe.png',
    9: 'ubw_4_2_9_risk.png',
    10: 'ubw_4_2_10_notifications.png',
    11: 'ubw_4_2_11_profile.png',
    12: 'ubw_4_2_12_admin_dashboard.png',
    13: 'ubw_4_2_13_admin_incidents.png',
    14: 'ubw_4_2_14_admin_masters.png',
}

def para_has_image(p):
    return ('graphic' in p._p.xml) or ('blip' in p._p.xml)

def embedded_size(p, doc):
    """Return (w,h) of first embedded image in paragraph, else None."""
    for blip in p._p.findall('.//' + qn('a:blip')):
        rid = blip.get(qn('r:embed'))
        if rid in doc.part.rels:
            try:
                blob = doc.part.rels[rid].target_part.blob
                im = Image.open(io.BytesIO(blob))
                return im.size
            except Exception:
                return None
    return None

def replace_picture(p, img_path, width_cm):
    for el in p._p.findall(qn('w:r')):
        if el.find(qn('w:drawing')) is not None:
            p._p.remove(el)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(0.2)
    p.paragraph_format.space_after = Cm(0.2)
    run = p.add_run()
    run.add_picture(img_path, width=Cm(width_cm))

doc = Document(DOC)
paras = doc.paragraphs
replaced = 0
for i, p in enumerate(paras):
    t = p.text.strip()
    mf = re.match(r'^Flowchart 4\.2\.(\d+):', t)
    if mf and i > 0 and para_has_image(paras[i - 1]):
        num = int(mf.group(1))
        sz = embedded_size(paras[i - 1], doc)
        # matplotlib flowcharts are 1063x1208 -> replace with user B&W, BIGGER
        if sz == (1063, 1208) and num in FLOW_MAP:
            replace_picture(paras[i - 1], os.path.join(UBW, FLOW_MAP[num]), 15.0)
            replaced += 1
            print('replaced Flowchart 4.2.%d (was %s)' % (num, sz))
    if t.startswith('Figure 4.1:') and i > 0 and para_has_image(paras[i - 1]):
        sz = embedded_size(paras[i - 1], doc)
        if sz == (1435, 1302):
            replace_picture(paras[i - 1], os.path.join(UBW, 'ubw_4_1_architecture.png'), 16.5)
            replaced += 1
            print('replaced Figure 4.1 architecture (was %s)' % (sz,))
    if t.startswith('Figure 4.3:') and i > 0 and para_has_image(paras[i - 1]):
        sz = embedded_size(paras[i - 1], doc)
        if sz == (1435, 1210):
            replace_picture(paras[i - 1], os.path.join(UBW, 'ubw_4_3_deployment.png'), 16.0)
            replaced += 1
            print('replaced Figure 4.3 deployment (was %s)' % (sz,))

print('total replaced:', replaced, '| inline images:', len(doc.inline_shapes))
doc.save(DOC)
print('SAVED:', DOC, round(os.path.getsize(DOC) / 1024), 'KB')